"""
Exportación de las cifras a Excel.

Entrega el mismo cálculo del informe en formato tabular, para quien necesite
cruzar, filtrar o construir sus propias tablas dinámicas. Tres hojas: por
municipio, por EAT y por subregión, más el diccionario de campos.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd

from . import config as cfg


def _programa(educacion: dict, nombre: str) -> dict:
    for fila in educacion["programas"]:
        if fila["nombre"] == nombre:
            return fila
    return {"total": 0, "conectadas": 0, "sin_dato": 0}


def _zona(educacion: dict, nombre: str) -> int:
    for fila in educacion.get("zonas", []):
        if fila["nombre"] == nombre:
            return fila["total"]
    return 0


def _fila_comun(o: dict) -> dict:
    edu, inf, sal = o["educacion"], o["infraestructura"], o["salud"]
    gob = _programa(edu, "Gobernación")
    mtc = _programa(edu, "MinTIC")
    sinp = _programa(edu, "Sin programa")
    return {
        "Índice de avance": round(o["indice"], 4),
        "Sedes educativas": edu["total"],
        "Sedes conectadas": edu["conectadas"],
        "Sedes sin dato": edu["sin_dato"],
        "% avance educativo": round(edu["avance"], 4) if edu["avance"] else None,
        "Sedes sin georreferenciar": edu["sin_geo"],
        "Sedes Gobernación": gob["total"],
        "Sedes Gobernación conectadas": gob["conectadas"],
        "Sedes MinTIC": mtc["total"],
        "Sedes MinTIC conectadas": mtc["conectadas"],
        "Sedes sin programa": sinp["total"],
        "Sedes rurales (Gobernación)": _zona(edu, "Rural"),
        "Sedes urbanas (Gobernación)": _zona(edu, "Urbana"),
        "Sedes sin clasificar por zona": _zona(edu, "Sin clasificar"),
        "DataCenter Gobernación": inf["datacenter_gobernacion"],
        "DataCenter propio": inf["datacenter_propio"],
        "Diagnóstico aplicado": inf["diagnosticados"],
        "Unidades entregadas (suma)": inf["equipos_total"],
        "Tipos de dispositivo": inf["equipos_tipos"],
        "AP internos": inf["ap_internos"],
        "AP externos": inf["ap_externos"],
        "AP sin funcionar": inf["ap_sin_funcionar"],
        "Puntos de conectividad": inf["zonas_wifi"],
        "Entornos educativos seguros": o["seguridad"]["puntos"],
        "Metros de fibra (proyecto dptal.)": o["fibra"]["metros"],
        "Municipios en proyecto de fibra": o["fibra"]["municipios"],
        "Puntos de salud": sal["puntos"],
        "Salud · kit satelital fijo": sal["kit_fijo"],
        "Salud · unidad itinerante": sal["kit_itinerante"],
        "Salud · sin kit satelital": sal["sin_kit"],
        "Personas formadas": o["apropiacion"]["personas"],
        "Personas formadas · comunidad": o["apropiacion"]["comunidad"],
        "Personas formadas · funcionarios": o["apropiacion"]["funcionarios"],
        "Actividades de formación": o["apropiacion"]["actividades"],
        "Personas capacitadas (acompañamiento)": o["apropiacion"]["acompanados"],
        "Sesiones de acompañamiento": o["apropiacion"]["sesiones"],
    }


def _tabla_municipios(datos: dict) -> pd.DataFrame:
    filas = []
    for m in datos["municipios"]:
        fila = {
            "COD_DANE_MPIO": m["cod"],
            "Municipio": m["nombre"],
            "Subregión": m["subregion"],
            "EAT": cfg.EAT_DISPLAY.get(m["eat"], m["eat"]),
            "Certificado en educación": "Sí" if m["certificado"] else "No",
        }
        fila.update(_fila_comun(m))
        fila.update({
            "Estado PETI": cfg.ETIQUETA_ESTADO.get(
                m["gobierno"]["estado_peti"], "Sin dato"),
            "Estado PAMUDA": cfg.ETIQUETA_ESTADO.get(
                m["gobierno"]["estado_pamuda"], "Sin dato"),
            "Ficha de madurez": m["gobierno"]["madurez"],
        })
        for eje, nombre, _, _ in cfg.EJES:
            valor = m["ejes"][eje]
            fila[f"{eje} · {nombre}"] = (
                None if valor is None else round(valor, 4)
            )
        filas.append(fila)
    return pd.DataFrame(filas).sort_values("Municipio")


def _tabla_agregado(datos: dict, clave: str, etiqueta: str) -> pd.DataFrame:
    filas = []
    for nombre, o in datos[clave].items():
        mostrado = cfg.EAT_DISPLAY.get(nombre, nombre)
        fila = {etiqueta: mostrado, "Municipios": o["municipios"]}
        fila.update(_fila_comun(o))
        gob = o.get("gobierno", {})
        for campo, titulo in (("peti", "PETI"), ("pamuda", "PAMUDA")):
            for item in gob.get(campo, []):
                nombre_estado = cfg.ETIQUETA_ESTADO.get(
                    item["nombre"], item["nombre"])
                fila[f"{titulo} · {nombre_estado}"] = item["total"]
        filas.append(fila)
    return pd.DataFrame(filas).sort_values(etiqueta)


def _diccionario() -> pd.DataFrame:
    return pd.DataFrame([
        ("COD_DANE_MPIO", "Código DANE del municipio, cinco dígitos como texto.",
         "Hoja Municipios"),
        ("Sedes educativas", "Total de sedes registradas en el municipio. Es el "
         "punto físico donde se instala la conectividad.",
         "SEDES EDUCATIVAS · CODIGO_DANE"),
        ("Sedes conectadas", "Sedes cuyo campo PROYECTO trae una tecnología. "
         "Mide completitud del inventario, no servicio activo.",
         "SEDES EDUCATIVAS"),
        ("% avance educativo", "Sedes conectadas sobre sedes totales.",
         "Calculado"),
        ("Puntos de conectividad", "Antes llamado zonas wifi. Puntos de acceso "
         "público georreferenciados.",
         "Consolidados + D. Gestión Territorial TIC"),
        ("Entornos educativos seguros", "Cada registro es una placa deportiva "
         "intervenida con cámaras e iluminación, NO una cámara individual.",
         "SEGURIDAD"),
        ("Sedes rurales / urbanas", "La sábana solo distingue zona en el "
         "programa de la Gobernación. Las demás sedes quedan sin clasificar.",
         "SEDES EDUCATIVAS · Observacion"),
        ("Unidades entregadas (suma)", "Sumatoria de la columna CANTIDAD. Suma "
         "unidades de distinta naturaleza: un DataCenter pesa igual que un "
         "inyector PoE. Léala junto a Tipos de dispositivo.",
         "INVENTARIO GENERAL"),
        ("AP internos / externos / sin funcionar", "Solo existen para los "
         "municipios con diagnóstico aplicado. En el resto son cero por falta "
         "de medición, no por ausencia de equipos.",
         "Diagnosticos DataCenters"),
        ("Salud · kit satelital fijo", "Sala de telemedicina del hospital.",
         "SALUD · token KIT_SATELITAL_STARLINK"),
        ("Salud · unidad itinerante", "Dotación portátil: maletín, batería, "
         "tablet y antena Starlink Mini.",
         "SALUD · token KIT_SATELITAL_ITINERANTE"),
        ("Metros de fibra", "Proyecto departamental de fibra óptica. Es "
         "independiente de las sedes conectadas por fibra de otros operadores.",
         "SEGURIDAD"),
        ("Estado PETI / PAMUDA", "Columnas ESTADO PETI y ESTADO PAMUDA.",
         "PETI-PAMUDA"),
        ("E1 a E7", "Ejes del índice, entre 0 y 1. E1 queda vacío en los "
         "municipios certificados en educación.",
         "Calculado"),
    ], columns=["Campo", "Qué significa", "Origen"])


def exportar(datos: dict, destino: Path | None = None) -> Path:
    destino = destino or (cfg.DIR_SALIDA / "cifras_tic_antioquia.xlsx")
    destino.parent.mkdir(parents=True, exist_ok=True)

    hojas = {
        "Por municipio": _tabla_municipios(datos),
        "Por EAT": _tabla_agregado(datos, "eats", "EAT"),
        "Por subregión": _tabla_agregado(datos, "subregiones", "Subregión"),
        "Diccionario": _diccionario(),
    }
    with pd.ExcelWriter(destino, engine="openpyxl") as writer:
        for nombre, df in hojas.items():
            df.to_excel(writer, sheet_name=nombre, index=False)

    _dar_formato(destino)
    return destino


def _dar_formato(ruta: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = load_workbook(ruta)
    verde = PatternFill("solid", fgColor="1B4D3E")
    for hoja in libro.worksheets:
        encabezados = [c.value for c in hoja[1]]
        for celda in hoja[1]:
            celda.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            celda.fill = verde
            celda.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
        hoja.row_dimensions[1].height = 42
        hoja.freeze_panes = "C2" if hoja.title != "Diccionario" else "A2"
        if hoja.max_row > 1:
            hoja.auto_filter.ref = (
                f"A1:{get_column_letter(hoja.max_column)}{hoja.max_row}"
            )
        for indice, nombre in enumerate(encabezados, start=1):
            letra = get_column_letter(indice)
            ancho = 14
            if nombre in ("Municipio", "Subregión", "EAT"):
                ancho = 26
            elif nombre == "COD_DANE_MPIO":
                ancho = 14
                for celda in hoja[letra][1:]:
                    celda.number_format = "@"
                    celda.alignment = Alignment(horizontal="center")
            elif nombre in ("Campo", "Qué significa", "Origen"):
                ancho = 46
                for celda in hoja[letra][1:]:
                    celda.alignment = Alignment(wrap_text=True, vertical="top")
            elif nombre and str(nombre).startswith("%"):
                for celda in hoja[letra][1:]:
                    celda.number_format = "0.0%"
            elif nombre and str(nombre).startswith("E") and "·" in str(nombre):
                for celda in hoja[letra][1:]:
                    celda.number_format = "0%"
            elif nombre == "Índice de avance":
                for celda in hoja[letra][1:]:
                    celda.number_format = "0.0%"
            hoja.column_dimensions[letra].width = ancho
    libro.save(ruta)
